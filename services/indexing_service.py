import csv
import os
import re
import uuid
from pathlib import Path
from typing import Any, Dict

import PyPDF2
from docx import Document

from app.core.config import settings
from app.core.logger import get_logger
from app.services.ollama_graphrag_embeddings import (
    graphrag_operation_lock,
    register_graphrag_ollama_embeddings,
)

logger = get_logger(service="indexing")


class IndexingService:
    """Prepare uploaded documents and build an isolated GraphRAG index."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}

    def __init__(self):
        self.project_dir = Path(settings.GRAPHRAG_PROJECT_DIR).expanduser().resolve()
        self.data_dir = self.project_dir / settings.GRAPHRAG_DATA_DIR
        self.config_path = self.data_dir / "settings.yaml"
        numba_cache_dir = self.data_dir / "cache" / "_numba"
        numba_cache_dir.mkdir(parents=True, exist_ok=True)
        os.environ["NUMBA_CACHE_DIR"] = str(numba_cache_dir)

    def _validate_configuration(self) -> None:
        if not self.project_dir.is_dir():
            raise FileNotFoundError(f"GraphRAG project directory does not exist: {self.project_dir}")
        if not self.data_dir.is_dir():
            raise FileNotFoundError(f"GraphRAG data directory does not exist: {self.data_dir}")
        if not self.config_path.is_file():
            raise FileNotFoundError(f"GraphRAG config file does not exist: {self.config_path}")

    @staticmethod
    def _read_text_file(file_path: Path) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return file_path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode text file: {file_path.name}")

    @staticmethod
    def _read_pdf(file_path: Path) -> str:
        pages = []
        with file_path.open("rb") as source:
            reader = PyPDF2.PdfReader(source)
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"\n--- Page {page_number} ---\n{text}")
        return "\n".join(pages)

    @staticmethod
    def _read_docx(file_path: Path) -> str:
        document = Document(file_path)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append(" | ".join(values))
        return "\n".join(blocks)

    @staticmethod
    def _read_csv(file_path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                with file_path.open("r", encoding=encoding, newline="") as source:
                    rows = list(csv.reader(source))
                return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode CSV file: {file_path.name}")

    def _extract_text(self, file_path: Path) -> str:
        extension = file_path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(self.SUPPORTED_EXTENSIONS))
            raise ValueError(f"Unsupported document type '{extension}'. Supported types: {supported}")

        if extension == ".pdf":
            text = self._read_pdf(file_path)
        elif extension == ".docx":
            text = self._read_docx(file_path)
        elif extension == ".csv":
            text = self._read_csv(file_path)
        else:
            text = self._read_text_file(file_path)

        if not text.strip():
            raise ValueError(
                "No text could be extracted from the document. "
                "Scanned PDFs require OCR before they can be indexed."
            )
        return text

    def _prepare_index_directories(self, user_id: int) -> tuple[str, Path, Path, Path, Path]:
        user_uuid = uuid.uuid5(uuid.NAMESPACE_DNS, f"user_{user_id}").hex
        index_id = f"{user_uuid}_{uuid.uuid4().hex}"
        input_dir = self.data_dir / "input" / index_id
        output_dir = self.data_dir / "output" / index_id
        cache_dir = self.data_dir / "cache" / index_id
        log_dir = self.data_dir / "logs" / index_id
        for directory in (input_dir, output_dir, cache_dir, log_dir):
            directory.mkdir(parents=True, exist_ok=True)
        return index_id, input_dir, output_dir, cache_dir, log_dir

    @staticmethod
    def _model_config_overrides() -> Dict[str, Any]:
        """Use the application's verified providers instead of a second secret file."""
        return {
            "models.default_chat_model.api_base": settings.DEEPSEEK_BASE_URL,
            "models.default_chat_model.api_key": settings.DEEPSEEK_API_KEY,
            "models.default_chat_model.model": settings.DEEPSEEK_MODEL,
            "models.default_chat_model.request_timeout": 60.0,
            "models.default_chat_model.max_retries": 2,
            "models.default_chat_model.concurrent_requests": 4,
            "models.default_embedding_model.api_base": (
                f"{settings.OLLAMA_BASE_URL.rstrip('/')}/v1"
            ),
            "models.default_embedding_model.api_key": "ollama",
            "models.default_embedding_model.model": settings.OLLAMA_EMBEDDING_MODEL,
            # Ollama on Windows can intermittently return 502 when GraphRAG
            # submits several embedding jobs at the same time.
            "models.default_embedding_model.request_timeout": 120.0,
            "models.default_embedding_model.max_retries": 4,
            "models.default_embedding_model.concurrent_requests": 1,
        }

    async def process_file(self, file_info: Dict[str, Any]) -> Dict[str, Any]:
        file_path = Path(file_info.get("path", ""))
        try:
            self._validate_configuration()
            if not file_path.is_file():
                raise FileNotFoundError(f"Uploaded file does not exist: {file_path}")

            # Import large native dependencies only for an indexing request.
            # Loading torch before spaCy/Thinc avoids the Windows c10.dll conflict.
            import torch  # noqa: F401
            import graphrag.api as api
            from graphrag.config.enums import IndexingMethod
            from graphrag.config.load_config import load_config
            from graphrag.logger.null_progress import NullProgressLogger

            user_id = int(file_info.get("user_id", 0))
            extracted_text = self._extract_text(file_path)
            index_id, input_dir, output_dir, cache_dir, log_dir = self._prepare_index_directories(user_id)

            text_filename = f"{file_path.stem}.txt"
            input_file_path = input_dir / text_filename
            input_file_path.write_text(extracted_text, encoding="utf-8")

            config_overrides = {
                **self._model_config_overrides(),
                "input.type": "file",
                "input.file_type": "text",
                "input.base_dir": str(input_dir),
                "input.file_pattern": rf".*{re.escape(text_filename)}$",
                "output.base_dir": str(output_dir),
                "cache.base_dir": str(cache_dir),
                "reporting.base_dir": str(log_dir),
                "vector_store.default_vector_store.db_uri": str(output_dir / "lancedb"),
            }
            # ModelManager is a process-wide singleton. Serialize GraphRAG work
            # and recreate its clients so every upload uses this index's cache.
            async with graphrag_operation_lock:
                register_graphrag_ollama_embeddings(
                    settings.OLLAMA_BASE_URL,
                    settings.OLLAMA_EMBEDDING_MODEL,
                )
                graphrag_config = load_config(self.data_dir, self.config_path, config_overrides)
                # RichProgressLogger emits emoji that fail on some Windows GBK consoles.
                progress_logger = NullProgressLogger()

                logger.info(f"Building GraphRAG index {index_id} for {file_path.name}")
                index_results = await api.build_index(
                    config=graphrag_config,
                    method=IndexingMethod.Standard,
                    is_update_run=False,
                    memory_profile=False,
                    progress_logger=progress_logger,
                )

            errors = [str(error) for result in index_results for error in (result.errors or [])]
            if errors:
                raise RuntimeError("; ".join(errors))

            required_outputs = (
                "entities.parquet",
                "communities.parquet",
                "community_reports.parquet",
                "relationships.parquet",
                "text_units.parquet",
            )
            missing_outputs = [name for name in required_outputs if not (output_dir / name).is_file()]
            if missing_outputs:
                raise RuntimeError(f"GraphRAG index is incomplete; missing: {', '.join(missing_outputs)}")

            return {
                "status": "success",
                "index_id": index_id,
                "user_id": user_id,
                "source_file": str(file_path),
                "input_file": str(input_file_path),
                "output_dir": str(output_dir),
                "extracted_characters": len(extracted_text),
            }
        except Exception as exc:
            logger.error(f"Document indexing failed: {exc}", exc_info=True)
            return {
                "status": "error",
                "file_path": str(file_path),
                "error": str(exc),
            }
